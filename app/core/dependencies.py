# app/core/dependencies.py
from __future__ import annotations
import os, re
from pathlib import Path
from typing import List, Dict, Any
from fastapi import HTTPException

# --- Loaders ---
import fitz  # PyMuPDF
from docx import Document as DocxDocument
try:
    from pptx import Presentation
except Exception:
    Presentation = None
try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

# --- Vector store factory ---
try:
    from app.vector.factory import get_vector_store
except ImportError as e:
    raise RuntimeError(
        "Vector factory not found. Ensure app/vector/factory.py exists and exports get_vector_store(session_id)."
    ) from e


# -------------------------
# Paths / guards
# -------------------------
def _session_dir(session_id: str) -> Path:
    return Path("./Storage") / session_id

def _files_dir(session_id: str) -> Path:
    return _session_dir(session_id) / "files"

def check_uploaded_files_exist(session_id: str, rag: Any | None = None) -> None:
    fdir = _files_dir(session_id)
    if not fdir.exists() or not any(p.is_file() for p in fdir.iterdir()):
        raise HTTPException(status_code=400, detail="Please upload files first.")


# -------------------------
# Text extractors
# -------------------------
def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")

def _read_docx(path: Path) -> str:
    doc = DocxDocument(path)
    return "\n".join(p.text for p in doc.paragraphs)

def _read_pdf(path: Path) -> str:
    out = []
    with fitz.open(path) as doc:
        for page in doc:
            out.append(page.get_text("text"))
    return "\n".join(out)

def _read_pptx(path: Path) -> str:
    if Presentation is None:
        return ""
    prs = Presentation(str(path))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text)
    return "\n".join(texts)

def _read_xlsx(path: Path) -> str:
    if load_workbook is None:
        return ""
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                lines.append(" | ".join(vals))
    return "\n".join(lines)

def _extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".txt", ".csv", ".json", ".xml"}:
        return _read_txt(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".pptx":
        return _read_pptx(path)
    if ext in {".xls", ".xlsx"}:
        return _read_xlsx(path)
    return ""


def _chunk(text: str, size: int = 900, overlap: int = 150) -> List[str]:
    if not text:
        return []
    chunks, i, n = [], 0, len(text)
    step = max(1, size - overlap)
    while i < n:
        chunks.append(text[i:i+size])
        i += step
    return chunks


# -------------------------
# Query quality helpers
# -------------------------
def _looks_nonsense(q: str) -> bool:
    s = (q or "").strip()
    if len(s) < 2:
        return True
    uniq = set(s)
    if len(uniq) <= 2 and any(ch.isalpha() for ch in uniq):
        return True
    return False

def _normalize_arabic(text: str) -> str:
    t = text or ""
    t = re.sub(r"[ًٌٍَُِّْـٰ]", "", t)     # diacritics
    t = re.sub(r"[أإآا]", "ا", t)          # unify alef
    t = re.sub(r"\s+", " ", t).strip()
    return t


# -------------------------
# Simple RAG wrapper
# -------------------------
class SimpleRAG:
    def __init__(self, session_id: str):
        self.session_id = session_id
        self.store = get_vector_store(session_id)

    def build_index(self, force_rebuild: bool = False):
        """
        Build/refresh per-session index from files in ./Storage/{session_id}/files.

        Strategy:
          - If force_rebuild=True → reset & rebuild.
          - Else → idempotent 'upsert' via stable IDs (filename:chunk_index) to avoid duplicates.
            (Relies on your FaissStore.add handling existing ids or full rebuild fallback.)
        """
        files_dir = _files_dir(self.session_id)
        if not files_dir.exists():
            return

        if force_rebuild:
            try:
                self.store.reset()
            except Exception:
                pass

        for p in files_dir.iterdir():
            if not p.is_file():
                continue
            text = _extract_text(p)
            if not text.strip():
                continue
            chunks = _chunk(text)
            if not chunks:
                continue

            metas = [{"filename": p.name, "chunk_index": i} for i in range(len(chunks))]
            ids   = [f"{p.name}:{i}" for i in range(len(chunks))]
            # Store should upsert on existing ids; if it doesn't, a rebuild (force_rebuild) is safer.
            self.store.add(chunks, metas, ids)

    def get_collection_stats(self) -> Dict[str, Any]:
        try:
            return self.store.stats()
        except Exception:
            return {"points_count": 0}

    def query(self, question: str) -> Dict[str, Any]:
        q_raw = (question or "").strip()
        if _looks_nonsense(q_raw):
            return {"answer": "Please write a clearer question related to your uploaded files.", "sources": []}

        # Normalize Arabic (English unaffected)
        q = _normalize_arabic(q_raw)

        top_k = int(os.getenv("TOP_K", 12))
        sim_threshold = float(os.getenv("SIM_THRESHOLD", "0.32"))

        # 1) Dense retrieval
        hits = self.store.search(q, top_k=top_k) or []
        if not hits:
            return {"answer": "I couldn't find a confident answer in your documents.", "sources": []}

        # 2) Similarity threshold
        strong = [h for h in hits if (h.get("similarity_score") or 0.0) >= sim_threshold]
        if not strong:
            return {"answer": "No passages were sufficiently related to your question.", "sources": []}

        # 3) Optional rerank (won't crash if module missing)
        try:
            from app.core.rerank import cross_encoder_rerank
            strong = cross_encoder_rerank(q, strong)
        except Exception:
            pass

        # 4) Extractive synthesis (safe import)
        chosen = strong[: min(8, len(strong))]
        try:
            from app.core.extractive import extractive_summarize
            answer_text = extractive_summarize(q, chosen)
        except Exception:
            answer_text = (chosen[0].get("content") or "")[:800]

        # 5) Optional local LLM polish via Ollama (safe import)
        if os.getenv("GENERATE_WITH_LLM", "false").lower() == "true":
            try:
                from app.core.local_llm import local_llm_grounded_synthesis
                answer_text = local_llm_grounded_synthesis(q, chosen, answer_text)
            except Exception:
                pass

        sources = [
            {
                "filename": (h.get("metadata") or {}).get("filename"),
                "chunk_index": (h.get("metadata") or {}).get("chunk_index"),
                "score": h.get("similarity_score") or h.get("rerank_score"),
            }
            for h in chosen
        ]
        return {"answer": answer_text, "sources": sources}


# Factory used by routes
def get_rag_system(session_id: str) -> SimpleRAG:
    return SimpleRAG(session_id=session_id)
